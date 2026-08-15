import os
import json
import logging

def get_file_size(file_path: str) -> int:
    try:
        return os.path.getsize(file_path)
    except Exception:
        return 0

def detect_document_type(file_path: str) -> str:
    """Detect file type by extension.
    Returns: 'pdf' | 'pptx' | 'xlsx' | 'csv' | 'code' | 'image' | 'unknown'"""
    ext = os.path.splitext(file_path)[1].lower()
    
    code_extensions = [
        '.py', '.js', '.jsx', '.ts', '.tsx', '.html', '.css', 
        '.c', '.cpp', '.h', '.java', '.go', '.rs', '.php', '.rb', 
        '.sh', '.json', '.yaml', '.yml', '.md'
    ]
    image_extensions = ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.svg', '.webp']
    
    if ext == '.pdf':
        return 'pdf'
    elif ext == '.pptx':
        return 'pptx'
    elif ext == '.xlsx':
        return 'xlsx'
    elif ext == '.csv':
        return 'csv'
    elif ext in code_extensions:
        return 'code'
    elif ext in image_extensions:
        return 'image'
    else:
        return 'unknown'

def read_pdf(file_path: str) -> dict:
    """Extract text from PDF using pdfplumber.
    Returns: {'status': 'success'|'error', 'text': str, 'pages': list[dict], 'page_count': int, 'metadata': dict}"""
    if not os.path.exists(file_path):
        return {'status': 'error', 'message': f"Oops! I couldn't find the file at {file_path}. Please check the path and try again."}
        
    try:
        import pdfplumber
    except ImportError:
        return {'status': 'error', 'message': "I need the 'pdfplumber' library to read PDFs. Please install it by running: pip install pdfplumber"}
        
    try:
        with pdfplumber.open(file_path) as pdf:
            pages = []
            full_text = []
            
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                pages.append({
                    'page_number': i + 1,
                    'text': text
                })
                full_text.append(text)
                
            return {
                'status': 'success',
                'text': '\n\n'.join(full_text),
                'pages': pages,
                'page_count': len(pdf.pages),
                'metadata': pdf.metadata if hasattr(pdf, 'metadata') else {}
            }
    except PermissionError:
        return {'status': 'error', 'message': f"I don't have permission to read the file at {file_path}. Make sure it's not open in another program."}
    except Exception as e:
        return {'status': 'error', 'message': f"Something went wrong while trying to read the PDF. The file might be corrupted or password-protected. Details: {str(e)}"}

def read_pptx(file_path: str) -> dict:
    """Extract slides from PowerPoint using python-pptx.
    Returns: {'status': 'success'|'error', 'slides': [{'index': int, 'title': str, 'content': str, 'notes': str}], 'slide_count': int}"""
    if not os.path.exists(file_path):
        return {'status': 'error', 'message': f"Oops! I couldn't find the file at {file_path}. Please check the path."}
        
    try:
        from pptx import Presentation
    except ImportError:
        return {'status': 'error', 'message': "I need the 'python-pptx' library to read PowerPoint files. Please install it by running: pip install python-pptx"}
        
    try:
        prs = Presentation(file_path)
        slides_data = []
        
        for i, slide in enumerate(prs.slides):
            title = ""
            if slide.shapes.title:
                title = slide.shapes.title.text
                
            content_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape != slide.shapes.title:
                    content_texts.append(shape.text)
                    
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text
                
            slides_data.append({
                'index': i + 1,
                'title': title,
                'content': '\n'.join(content_texts),
                'notes': notes
            })
            
        return {
            'status': 'success',
            'slides': slides_data,
            'slide_count': len(prs.slides)
        }
    except PermissionError:
        return {'status': 'error', 'message': f"Permission denied for {file_path}. Please close the file if it's open in PowerPoint."}
    except Exception as e:
        return {'status': 'error', 'message': f"Failed to read the PowerPoint presentation. It might be corrupted. Details: {str(e)}"}

def read_excel(file_path: str) -> dict:
    """Read spreadsheet data using openpyxl.
    Returns: {'status': 'success'|'error', 'sheets': [{'name': str, 'headers': list, 'rows': list[list], 'row_count': int}]}"""
    if not os.path.exists(file_path):
        return {'status': 'error', 'message': f"Oops! I couldn't find the file at {file_path}."}
        
    try:
        import openpyxl
    except ImportError:
        return {'status': 'error', 'message': "I need the 'openpyxl' library to read Excel files. Please install it by running: pip install openpyxl"}
        
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets_data = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            
            headers = []
            data_rows = []
            
            if rows:
                headers = list(rows[0]) if rows[0] else []
                data_rows = [list(r) for r in rows[1:] if any(cell is not None for cell in r)]
                
            sheets_data.append({
                'name': sheet_name,
                'headers': headers,
                'rows': data_rows,
                'row_count': len(data_rows)
            })
            
        return {
            'status': 'success',
            'sheets': sheets_data
        }
    except PermissionError:
        return {'status': 'error', 'message': f"Permission denied for {file_path}. Please make sure it's not open in Excel."}
    except Exception as e:
        return {'status': 'error', 'message': f"Could not read the Excel file. Details: {str(e)}"}

def read_code_file(file_path: str) -> dict:
    """Read a source code file with language detection.
    Returns: {'status': 'success'|'error', 'content': str, 'language': str, 'line_count': int, 'file_size': int}"""
    if not os.path.exists(file_path):
        return {'status': 'error', 'message': f"Oops! I couldn't find the file at {file_path}."}
        
    ext = os.path.splitext(file_path)[1].lower()
    
    language_map = {
        '.py': 'Python',
        '.js': 'JavaScript',
        '.jsx': 'React/JSX',
        '.ts': 'TypeScript',
        '.tsx': 'React/TSX',
        '.html': 'HTML',
        '.css': 'CSS',
        '.json': 'JSON',
        '.md': 'Markdown',
        '.csv': 'CSV'
    }
    language = language_map.get(ext, 'Plain Text')
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
            lines = content.splitlines()
            
        return {
            'status': 'success',
            'content': content,
            'language': language,
            'line_count': len(lines),
            'file_size': get_file_size(file_path)
        }
    except UnicodeDecodeError:
        return {'status': 'error', 'message': f"Could not read {file_path} as text. It might be a binary file."}
    except PermissionError:
        return {'status': 'error', 'message': f"Permission denied to read {file_path}."}
    except Exception as e:
        return {'status': 'error', 'message': f"An error occurred while reading the file: {str(e)}"}

def extract_document_summary(file_path: str) -> dict:
    """Auto-detect file type and extract a text summary suitable for LLM analysis.
    Returns: {'status': 'success'|'error', 'file_type': str, 'summary_text': str, 'metadata': dict}"""
    
    file_type = detect_document_type(file_path)
    
    if file_type == 'unknown' or file_type == 'image':
        return {
            'status': 'error',
            'message': f"I'm sorry, but I don't know how to extract a text summary from a file of type: {file_type}."
        }
        
    try:
        if file_type == 'pdf':
            result = read_pdf(file_path)
            if result.get('status') == 'success':
                return {
                    'status': 'success',
                    'file_type': file_type,
                    'summary_text': result.get('text', ''),
                    'metadata': {'page_count': result.get('page_count'), 'pdf_metadata': result.get('metadata')}
                }
            return result
            
        elif file_type == 'pptx':
            result = read_pptx(file_path)
            if result.get('status') == 'success':
                summary_lines = []
                for slide in result.get('slides', []):
                    summary_lines.append(f"Slide {slide['index']}: {slide['title']}")
                    if slide['content']:
                        summary_lines.append(slide['content'])
                    if slide['notes']:
                        summary_lines.append(f"Notes: {slide['notes']}")
                    summary_lines.append("---")
                
                return {
                    'status': 'success',
                    'file_type': file_type,
                    'summary_text': '\n'.join(summary_lines),
                    'metadata': {'slide_count': result.get('slide_count')}
                }
            return result
            
        elif file_type == 'xlsx':
            result = read_excel(file_path)
            if result.get('status') == 'success':
                summary_lines = []
                for sheet in result.get('sheets', []):
                    summary_lines.append(f"Sheet: {sheet['name']} ({sheet['row_count']} rows)")
                    if sheet['headers']:
                        summary_lines.append("Headers: " + ", ".join(str(h) for h in sheet['headers']))
                    # Include up to first 5 rows for the summary preview
                    preview_rows = sheet['rows'][:5]
                    for r in preview_rows:
                        summary_lines.append(", ".join(str(cell) for cell in r))
                    summary_lines.append("---")
                    
                return {
                    'status': 'success',
                    'file_type': file_type,
                    'summary_text': '\n'.join(summary_lines),
                    'metadata': {'sheet_count': len(result.get('sheets', []))}
                }
            return result
            
        elif file_type in ['csv', 'code']:
            result = read_code_file(file_path)
            if result.get('status') == 'success':
                return {
                    'status': 'success',
                    'file_type': file_type,
                    'summary_text': result.get('content', ''),
                    'metadata': {'language': result.get('language'), 'line_count': result.get('line_count')}
                }
            return result
            
    except Exception as e:
        return {'status': 'error', 'message': f"Unexpected error during summary extraction: {str(e)}"}
